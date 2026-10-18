#!/usr/bin/env python3
"""Build the Indonesian OpsMind deck from the official template."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from build_deck import (
    AMBER,
    GITHUB_URL,
    NAVY,
    OUT_DIR,
    ROOT,
    SLATE,
    TEAL,
    TEAM_LEAD,
    TEAM_NAME,
    VIDEO_URL,
    DEMO_URL,
    add_caption,
    delete_body,
    set_body,
    set_paragraphs,
    set_title,
    text_shapes,
)

TEMPLATE = ROOT / "Submission Template _ AI Builder Cup - Indonesian.pptx"
OUT_PPTX = OUT_DIR / "OpsMind_AI_Deck_ID.pptx"


def build() -> Path:
    prs = Presentation(TEMPLATE)
    slides = list(prs.slides)

    # 1 — detail tim
    shape = text_shapes(slides[0])[0]
    set_paragraphs(
        shape,
        [
            ("OpsMind — AI Operational Intelligence", 22, True, NAVY),
            ("Google Cloud AI Builder Cup 2026 · JAPAC", 11, False, TEAL),
            (f"Tim: {TEAM_NAME}   ·   Ketua: {TEAM_LEAD}", 11, True, SLATE),
            (
                "Masalah: data operasi terpisah; lonjakan listrik tidak selalu terlihat dari produksi.",
                10.5,
                False,
                SLATE,
            ),
            (
                "Prototipe: metrik deterministik + penjelasan Gemini. AI menyarankan; manusia memutuskan.",
                10.5,
                False,
                SLATE,
            ),
        ],
        space_after=3,
    )

    # 2 — ringkasan eksekutif
    set_body(
        slides[1],
        [
            ("Ringkasan eksekutif", 16, True, NAVY),
            ("Data operasi terpisah antara produksi, listrik, jam operasi dan cuaca.", 10.5, False, SLATE),
            ("OpsMind menghitung intensitas, baseline, deviasi dan estimasi biaya potensial.", 10.5, False, SLATE),
            ("Gemini hanya menjelaskan bukti yang diberikan dan menyebut data yang belum tersedia.", 10.5, False, SLATE),
            ("Model tidak menghitung angka. Semua data demo berlabel SIMULATED DATA.", 10.5, False, SLATE),
            ("AI menyarankan; manusia memutuskan.", 11, True, TEAL),
        ],
        left=5.0,
        top=1.45,
        width=4.55,
        height=3.6,
        size_default=10.5,
    )

    # 3 — ide
    set_title(slides[2], "Ringkasan ide")
    set_body(
        slides[2],
        [
            ("• OpsMind adalah prototipe AI Operational Intelligence untuk manufaktur kecil–menengah.", None, False, SLATE),
            ("• Menghitung intensitas energi (kWh/unit), baseline, deviasi dan estimasi biaya kelebihan.", None, False, SLATE),
            ("• Gemini menjelaskan bukti, memisahkan fakta dari hipotesis, dan menyarankan pemeriksaan manusia.", None, False, SLATE),
            ("• Semua data demo berlabel SIMULATED DATA. Tidak ada klaim penghematan.", None, False, SLATE),
        ],
        size_default=12,
    )

    # 4 — peluang / USP
    set_title(slides[3], "Peluang")
    set_body(
        slides[3],
        [
            ("• Diferensiasi: investigasi berbasis bukti, bukan chatbot umum atau alarm ambang batas.", None, False, SLATE),
            ("• Setiap observasi mengutip bukti E1–E13; setiap hipotesis mencantumkan bukti yang belum tersedia.", None, False, SLATE),
            ("• Panel biaya menampilkan rumus dan asumsi — bukan tagihan atau penghematan terverifikasi.", None, False, SLATE),
            ("• Cuaca adalah konteks, bukan bukti sebab-akibat; perubahan before/after tidak diklaim sebagai akibat tindakan.", None, False, SLATE),
            ("• USP: alur human-in-the-loop dari deteksi → investigasi → tindakan → verifikasi.", 11.5, True, TEAL),
        ],
        size_default=11.5,
    )

    # 5 — fitur
    set_title(slides[4], "Fitur")
    set_body(
        slides[4],
        [
            ("• Dasbor harian: produksi, listrik, intensitas energi dan estimasi biaya potensial.", None, False, SLATE),
            ("• Deteksi transparan: baseline, ambang heuristik 15% dan tingkat severity.", None, False, SLATE),
            ("• Investigasi AI: Gemini atau mock deterministik berlabel; output divalidasi Zod.", None, False, SLATE),
            ("• Bukti: 13 item (E1–E13), fakta vs hipotesis, data yang hilang dan keterbatasan.", None, False, SLATE),
            ("• Alur tindakan: Open / In progress / Completed dengan timestamp dan catatan.", None, False, SLATE),
            ("• Perbandingan before/after dengan peringatan “perubahan yang diamati saja”.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 6 — alur investigasi
    set_title(slides[5], "Alur investigasi")
    delete_body(slides[5])
    slides[5].shapes.add_picture(
        str(ROOT / "submission-assets/diagrams/process.png"), Inches(0.55), Inches(1.35), width=Inches(8.9)
    )
    add_caption(
        slides[5],
        "Mesin deterministik menghitung; Gemini menjelaskan; manusia memutuskan. Kegagalan model disimpan sebagai FAILED.",
        0.55,
        5.24,
        8.9,
    )

    # 7 — walkthrough
    set_title(slides[6], "Tampilan produk — Overview")
    slides[6].shapes.add_picture(
        str(ROOT / "submission-assets/screenshots/overview-card.png"), Inches(0.55), Inches(1.35), width=Inches(8.9)
    )
    add_caption(
        slides[6],
        "30 hari simulasi · 17 Sep 2026: +36,0% terhadap baseline · estimasi potensial Rp288.884.",
        0.55,
        5.24,
        8.9,
    )

    # 8 — arsitektur
    set_title(slides[7], "Arsitektur")
    delete_body(slides[7])
    slides[7].shapes.add_picture(
        str(ROOT / "submission-assets/diagrams/architecture.png"), Inches(0.55), Inches(1.35), width=Inches(8.9)
    )
    add_caption(
        slides[7],
        "Next.js → route handler → validasi → service → PostgreSQL; output Gemini divalidasi sebelum disimpan.",
        0.55,
        5.24,
        8.9,
    )

    # 9 — teknologi
    set_title(slides[8], "Teknologi")
    set_body(
        slides[8],
        [
            ("• Next.js 16 App Router · React 19 · TypeScript 7 strict · design token CSS kustom.", None, False, SLATE),
            ("• PostgreSQL + Prisma 7 driver adapter · kolom decimal dan date native.", None, False, SLATE),
            ("• Gemini API melalui REST server-side · validasi output Zod 4 · fallback mock berlabel.", None, False, SLATE),
            ("• Vitest 5 (79 test) · Playwright E2E · Docker · GitHub Actions.", None, False, SLATE),
            ("• PostgreSQL embedded untuk lokal · target deployment: Google Cloud Run.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 10 — biaya
    set_title(slides[9], "Estimasi biaya implementasi (opsional)")
    set_body(
        slides[9],
        [
            ("• Pengembangan: PostgreSQL embedded lokal + mock investigator; tanpa biaya cloud.", None, False, SLATE),
            ("• Deployment: Cloud Run scale-to-zero + PostgreSQL managed + token Gemini; biaya tergantung trafik dan region.", None, False, SLATE),
            ("• Aplikasi menampilkan estimasi potensial: kWh berlebih × tarif, beserta asumsinya.", None, False, SLATE),
            ("• Tidak ada klaim penghematan terverifikasi; diperlukan pengukuran before/after yang terkontrol.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 11 — tangkapan layar
    set_title(slides[10], "Tangkapan layar prototipe — data simulasi")
    delete_body(slides[10])
    cards = [
        ("overview-card.png", "Overview — KPI, tren dan antrean investigasi"),
        ("investigation-card.png", "Investigasi — bukti E1–E13 dan penjelasan AI"),
        ("actions-card.png", "Tindakan — status dan perbandingan before/after"),
        ("data-card.png", "Data — label sumber dan filter tanggal"),
    ]
    positions = [(0.55, 1.38), (4.95, 1.38), (0.55, 3.42), (4.95, 3.42)]
    for (filename, caption), (left, top) in zip(cards, positions):
        slides[10].shapes.add_picture(
            str(ROOT / "submission-assets/screenshots" / filename), Inches(left), Inches(top), width=Inches(3.9)
        )
        add_caption(slides[10], caption, left, top + 1.73, 3.9, size=8.5)

    # 12 — performa
    set_title(slides[11], "Performa dan verifikasi")
    set_body(
        slides[11],
        [
            ("• 79 test otomatis: mesin kalkulasi, validasi, kegagalan AI dan integrasi API.", None, False, SLATE),
            ("• 3 test Playwright E2E: overview, navigasi sidebar dan alur investigasi.", None, False, SLATE),
            ("• Typecheck dan build produksi lulus; dataset demo 30 hari deterministik.", None, False, SLATE),
            ("• Health check: /api/health melaporkan status database dan provider AI.", None, False, SLATE),
            ("• Tidak diklaim: throughput produksi, benchmark latency, atau penghematan nyata.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 13 — roadmap
    set_title(slides[12], "Roadmap")
    set_body(
        slides[12],
        [
            ("• Provider cuaca nyata dengan cakupan, ketentuan, atribusi dan evaluasi rate limit.", None, False, SLATE),
            ("• Upload CSV dengan preview/validasi; dukungan multi-site.", None, False, SLATE),
            ("• Autentikasi/RBAC dan rate limit di gateway.", None, False, SLATE),
            ("• Studi before/after terkontrol dengan penghematan terukur.", None, False, SLATE),
            ("• Evaluasi model pada data operasi nyata; versioning prompt/model.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 14 — tautan
    set_title(slides[13], "Tautan submission")
    links = [
        ("• Repositori GitHub publik: " + (GITHUB_URL or "lihat form submission Hack2skill."), None, False, SLATE),
        ("• Video demo (<3 menit): " + (VIDEO_URL or "lihat form submission."), None, False, SLATE),
        ("• Prototipe live: " + (DEMO_URL or "lihat form submission."), None, False, SLATE),
        ("• Health check: /api/health (status database + provider AI).", None, False, SLATE),
        ("• Form submission adalah sumber utama untuk URL final.", 10.5, True, TEAL),
    ]
    set_body(slides[13], links, size_default=11.5)

    # 15 — penutup
    box = slides[14].shapes.add_textbox(Inches(0.55), Inches(1.45), Inches(8.9), Inches(3.6))
    set_paragraphs(
        box,
        [
            ("OpsMind — AI Operational Intelligence", 24, True, NAVY),
            ("Tanyakan: apa buktinya?", 17, True, TEAL),
            (
                "Kandidat demo: 17 Sep 2026, +36,0% di atas baseline — estimasi potensial, bukan penghematan terverifikasi.",
                12,
                False,
                SLATE,
            ),
            ("AI menyarankan; manusia memutuskan.", 12, True, AMBER),
            ("Terima kasih.", 14, False, SLATE),
        ],
        space_after=8,
    )

    prs.core_properties.title = "OpsMind — AI Operational Intelligence"
    prs.core_properties.author = TEAM_NAME
    prs.core_properties.subject = "Google Cloud AI Builder Cup 2026 · JAPAC"
    prs.core_properties.keywords = "Gemini, Cloud Run, efisiensi energi, operasi, prototipe"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    return OUT_PPTX


if __name__ == "__main__":
    print(f"wrote {build()}")
