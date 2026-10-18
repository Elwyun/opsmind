#!/usr/bin/env python3
"""Build the OpsMind submission deck from the official AI Builder Cup template.

Run (PowerPoint/LibreOffice only needed for the PDF step):
    PYTHONPATH=./.deck-tools python scripts/build_deck.py

Optional environment variables fill the links slide:
    TEAM_NAME, TEAM_LEAD, GITHUB_URL, VIDEO_URL, DEMO_URL
"""

from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "Submission Template _ AI Builder Cup.pptx"
OUT_DIR = ROOT / "submission"
OUT_PPTX = OUT_DIR / "OpsMind_AI_Deck_EN.pptx"

NAVY = RGBColor(0x0E, 0x3A, 0x5C)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
TEAL = RGBColor(0x0D, 0x7D, 0x78)
AMBER = RGBColor(0xB4, 0x53, 0x09)

TEAM_NAME = os.environ.get("TEAM_NAME", "OpsMind AI")
TEAM_LEAD = os.environ.get("TEAM_LEAD", "Sisawaktu 11")
GITHUB_URL = os.environ.get("GITHUB_URL", "").strip()
VIDEO_URL = os.environ.get("VIDEO_URL", "").strip()
DEMO_URL = os.environ.get("DEMO_URL", "").strip()


def set_paragraphs(shape, specs, space_after: int = 5) -> None:
    """Replace a text frame with (text, size, bold, color) paragraphs."""
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, (text, size, bold, color) in enumerate(specs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 0
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(space_after)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def text_shapes(slide):
    return [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]


def set_title(slide, text: str) -> None:
    for shape in text_shapes(slide):
        if shape.top < Inches(1.0):
            set_paragraphs(shape, [(text, 17, True, NAVY)], space_after=0)
            return


def body_shape(slide):
    candidates = [shape for shape in text_shapes(slide) if shape.top > Inches(1.0)]
    return candidates[-1] if candidates else None


def set_body(slide, specs, *, left=0.55, top=1.35, width=8.9, height=3.9, size_default=None) -> None:
    shape = body_shape(slide)
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.left = Inches(left)
    shape.top = Inches(top)
    shape.width = Inches(width)
    shape.height = Inches(height)
    if size_default is not None:
        specs = [(text, size_default if size is None else size, bold, color) for text, size, bold, color in specs]
    set_paragraphs(shape, specs)


def delete_body(slide) -> None:
    for shape in list(text_shapes(slide)):
        if shape.top > Inches(1.0):
            shape._element.getparent().remove(shape._element)


def add_caption(slide, text: str, left: float, top: float, width: float, size: float = 9.5) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.3))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.margin_left = 0
    frame.margin_right = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = MUTED


def build() -> Path:
    prs = Presentation(TEMPLATE)
    slides = list(prs.slides)

    # 1 — team details
    slide = slides[0]
    shape = text_shapes(slide)[0]
    set_paragraphs(
        shape,
        [
            ("OpsMind — AI Operational Intelligence", 22, True, NAVY),
            ("Google Cloud AI Builder Cup 2026 · JAPAC", 11, False, TEAL),
            (f"Team: {TEAM_NAME}   ·   Team lead: {TEAM_LEAD}", 11, True, SLATE),
            (
                "Problem: operations data is split; a power spike does not obviously track production.",
                10.5,
                False,
                SLATE,
            ),
            (
                "Prototype: deterministic metrics + Gemini explanations. AI suggests; a person decides.",
                10.5,
                False,
                SLATE,
            ),
        ],
        space_after=3,
    )

    # 2 — executive summary (replaces the template note)
    slide = slides[1]
    set_body(
        slide,
        [
            ("Executive summary", 16, True, NAVY),
            ("Operations data is split across production, electricity, hours and weather.", 10.5, False, SLATE),
            ("OpsMind calculates intensity, a trailing baseline, deviation and a potential cost estimate.", 10.5, False, SLATE),
            ("Gemini explains only the supplied evidence and lists what is still missing.", 10.5, False, SLATE),
            ("The model never calculates the numbers. Every demo row is labelled SIMULATED DATA.", 10.5, False, SLATE),
            ("AI suggests; a person decides.", 11, True, TEAL),
        ],
        left=5.0,
        top=1.45,
        width=4.55,
        height=3.6,
        size_default=10.5,
    )

    # 3 — idea
    slide = slides[2]
    set_title(slide, "Brief about the idea")
    set_body(
        slide,
        [
            ("• OpsMind is an AI Operational Intelligence prototype for SME manufacturers.", None, False, SLATE),
            ("• It computes energy intensity (kWh/unit), a trailing baseline, deviation and a potential excess-cost estimate.", None, False, SLATE),
            ("• Gemini explains the evidence, separates facts from hypotheses and recommends human checks.", None, False, SLATE),
            ("• All demo data is labelled SIMULATED DATA. No saving is claimed.", None, False, SLATE),
        ],
        size_default=12,
    )

    # 4 — opportunities / USP
    slide = slides[3]
    set_title(slide, "Opportunities")
    set_body(
        slide,
        [
            ("• Differentiation: evidence-first investigation, not a generic chatbot or a single threshold alert.", None, False, SLATE),
            ("• Every observation cites evidence E1–E13; each hypothesis lists the evidence that is still missing.", None, False, SLATE),
            ("• The cost panel shows the formula and assumptions — not a claimed bill or verified saving.", None, False, SLATE),
            ("• Weather is context, not proof; before/after changes are reported without asserting causation.", None, False, SLATE),
            ("• USP: a human-in-the-loop path from detection → investigation → action → verification.", 11.5, True, TEAL),
        ],
        size_default=11.5,
    )

    # 5 — features
    slide = slides[4]
    set_title(slide, "Features")
    set_body(
        slide,
        [
            ("• Daily dashboard: production, electricity, energy intensity and potential additional cost.", None, False, SLATE),
            ("• Transparent detection: trailing baseline, 15% heuristic threshold and severity bands.", None, False, SLATE),
            ("• AI investigation: Gemini or a labelled deterministic mock; Zod-validated output.", None, False, SLATE),
            ("• Evidence: 13 items (E1–E13), facts vs hypotheses, missing inputs and limitations.", None, False, SLATE),
            ("• Action workflow: Open / In progress / Completed with timestamps and completion notes.", None, False, SLATE),
            ("• Before/after comparison with an explicit “observed change only” caveat.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 6 — process flow
    slide = slides[5]
    set_title(slide, "Investigation flow")
    delete_body(slide)
    slide.shapes.add_picture(str(ROOT / "submission-assets/diagrams/process.png"), Inches(0.55), Inches(1.35), width=Inches(8.9))
    add_caption(slide, "The deterministic engine calculates; Gemini explains; a person decides. Model failure is stored as FAILED.", 0.55, 5.24, 8.9)

    # 7 — product walkthrough
    slide = slides[6]
    set_title(slide, "Product walkthrough — Overview")
    slide.shapes.add_picture(str(ROOT / "submission-assets/screenshots/overview-card.png"), Inches(0.55), Inches(1.35), width=Inches(8.9))
    add_caption(slide, "30 simulated days · 17 Sep 2026: +36.0% vs the trailing baseline · potential estimate Rp288,884.", 0.55, 5.24, 8.9)

    # 8 — architecture
    slide = slides[7]
    set_title(slide, "Architecture")
    delete_body(slide)
    slide.shapes.add_picture(str(ROOT / "submission-assets/diagrams/architecture.png"), Inches(0.55), Inches(1.35), width=Inches(8.9))
    add_caption(slide, "Next.js → route handlers → validation → services → PostgreSQL; Gemini output is validated before storage.", 0.55, 5.24, 8.9)

    # 9 — technology
    slide = slides[8]
    set_title(slide, "Technology stack")
    set_body(
        slide,
        [
            ("• Next.js 16 App Router · React 19 · TypeScript 7 strict · custom CSS design tokens.", None, False, SLATE),
            ("• PostgreSQL + Prisma 7 driver adapter · native decimal and date columns.", None, False, SLATE),
            ("• Gemini API via server-side REST · Zod 4 output validation · labelled deterministic mock fallback.", None, False, SLATE),
            ("• Vitest 5 (79 tests) · Playwright E2E · Docker · GitHub Actions.", None, False, SLATE),
            ("• Local embedded PostgreSQL · deployment target: Google Cloud Run.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 10 — cost (optional, kept honest)
    slide = slides[9]
    set_title(slide, "Estimated implementation cost (optional)")
    set_body(
        slide,
        [
            ("• Development: local embedded PostgreSQL + labelled mock investigator; no cloud spend required.", None, False, SLATE),
            ("• Deployment: Cloud Run scale-to-zero + managed PostgreSQL + Gemini token usage; cost depends on traffic and region.", None, False, SLATE),
            ("• The app reports a potential estimate: excess kWh × tariff, with the tariff and assumptions shown.", None, False, SLATE),
            ("• It does not claim verified savings; a controlled before/after measurement is required.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 11 — snapshots
    slide = slides[10]
    set_title(slide, "Prototype snapshots — simulated data")
    delete_body(slide)
    cards = [
        ("overview-card.png", "Overview — KPIs, trend and investigation queue"),
        ("investigation-card.png", "Investigation — evidence E1–E13 and AI explanation"),
        ("actions-card.png", "Actions — open/completed and before/after comparison"),
        ("data-card.png", "Data — source labels and date filter"),
    ]
    positions = [(0.55, 1.38), (4.95, 1.38), (0.55, 3.42), (4.95, 3.42)]
    for (filename, caption), (left, top) in zip(cards, positions):
        slide.shapes.add_picture(
            str(ROOT / "submission-assets/screenshots" / filename),
            Inches(left),
            Inches(top),
            width=Inches(3.9),
        )
        add_caption(slide, caption, left, top + 1.73, 3.9, size=8.5)

    # 12 — performance and verification
    slide = slides[11]
    set_title(slide, "Performance and verification")
    set_body(
        slide,
        [
            ("• 79 automated tests: calculation engine, validation, AI failure modes and API integration.", None, False, SLATE),
            ("• 3 Playwright E2E tests: overview, sidebar navigation and the investigate journey.", None, False, SLATE),
            ("• Typecheck and production build pass; the 30-day demo dataset is deterministic.", None, False, SLATE),
            ("• Health check: /api/health reports database and AI-provider status.", None, False, SLATE),
            ("• Not claimed: production throughput, a latency benchmark, or real savings.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 13 — roadmap
    slide = slides[12]
    set_title(slide, "Roadmap")
    set_body(
        slide,
        [
            ("• Real weather provider with coverage, terms, attribution and rate-limit review.", None, False, SLATE),
            ("• CSV upload with preview/validation; multi-site support.", None, False, SLATE),
            ("• Authentication/RBAC and gateway-level rate limits.", None, False, SLATE),
            ("• Controlled before/after studies with measured savings.", None, False, SLATE),
            ("• Model evaluation on real operations data; prompt/model versioning.", None, False, SLATE),
        ],
        size_default=11.5,
    )

    # 14 — links (no invented URLs)
    slide = slides[13]
    set_title(slide, "Submission links")
    links = [
        ("• GitHub public repository: " + (GITHUB_URL or "see the Hack2skill submission form."), None, False, SLATE),
        ("• Demo video (<3 min): " + (VIDEO_URL or "see the submission form."), None, False, SLATE),
        ("• Live prototype: " + (DEMO_URL or "see the submission form."), None, False, SLATE),
        ("• Health check: /api/health (database + AI-provider status).", None, False, SLATE),
        ("• The submission form is the source of truth for the final URLs.", 10.5, True, TEAL),
    ]
    set_body(slide, links, size_default=11.5)

    # 15 — close
    slide = slides[14]
    box = slide.shapes.add_textbox(Inches(0.55), Inches(1.45), Inches(8.9), Inches(3.6))
    set_paragraphs(
        box,
        [
            ("OpsMind — AI Operational Intelligence", 24, True, NAVY),
            ("Ask: what is the evidence?", 17, True, TEAL),
            (
                "Demo candidate: 17 Sep 2026, +36.0% above the trailing baseline — a potential estimate, not a verified saving.",
                12,
                False,
                SLATE,
            ),
            ("AI suggests; a person decides.", 12, True, AMBER),
            ("Thank you.", 14, False, SLATE),
        ],
        space_after=8,
    )

    prs.core_properties.title = "OpsMind — AI Operational Intelligence"
    prs.core_properties.author = TEAM_NAME
    prs.core_properties.subject = "Google Cloud AI Builder Cup 2026 · JAPAC"
    prs.core_properties.keywords = "Gemini, Cloud Run, energy efficiency, operations, prototype"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    return OUT_PPTX


if __name__ == "__main__":
    output = build()
    print(f"wrote {output}")
